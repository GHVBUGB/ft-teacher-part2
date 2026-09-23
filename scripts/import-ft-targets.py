"""Import this reviewed FT workbook batch; never modifies the source workbook.

Run with a Python environment containing openpyxl:
    python scripts/import-ft-targets.py '/path/FT Target Words.xlsx'
Only catalog content and offline provenance are written. Grammar is preserved.
"""
import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
VERSION = 'ft-target-words-2026-09-14-v1'
SOURCE_SHA256 = '0039151d934f7d2f6b28c04192f0bc9493f992344ec06270834ee7a89f6b2b0a'

# Explicit expansion avoids treating slash notation as a literal reading target.
EXPANSIONS = {
    28: ["It's a brown bear.", "It's a brown monster."],
    29: ["It's a yellow bear.", "It's a yellow car."],
    31: ['I like the brown bear.', 'I like the yellow bear.'],
    32: ['I like the yellow car.', 'I like the green car.'],
    33: ['I like the green monster.', 'I like the brown monster.'],
    78: [f"I'm wearing {item}." for item in ['a shirt', 'shorts', 'sneakers', 'a coat', 'pants', 'boots']],
    82: ['You are good at running with the ball.', 'You are good at shooting.'],
    83: ['You are good at passing the ball.', 'You are good at taking away the ball.'],
    84: [f'I will choose {item} for my football player.' for item in ['football socks', 'a football shirt', 'football boots', 'football shorts']],
    90: [f"It's {item}." for item in ['sunny and hot', 'rainy and cool', 'snowy and cold']],
    91: [f"Let's go to the {item}." for item in ['beach', 'town', 'hills']],
    92: ['We can go swimming.', 'We can go surfing.'],
    93: ['We can go shopping.', 'We can ride a bike.'],
    94: ['We can make a snowman.', 'We can go skiing.'],
    95: ["I'd like a shawarma.", "I'd like a burger."],
    **{row: [f"I'd like some {item}.", f'I prefer some {item}.'] for row, item in enumerate(['tomatoes', 'beef', 'cheese', 'water', 'burger', 'bell peppers', 'chicken', 'sauce', 'juice'], 96) if row != 100},
}
CORRECTIONS = {
    'B124': ('achipelago', 'archipelago', '修正明显拼写错误'),
    'D42': ('They are bird.', 'They are birds.', '修正 are 后复数名词'),
    'D141': ("Let's try bumber cars.", "Let's try bumper cars.", '修正明显拼写错误'),
}


def clean(text):
    return re.sub(r'\s+', ' ', str(text or '')).strip()


def main(source):
    assert hashlib.sha256(source.read_bytes()).hexdigest() == SOURCE_SHA256, 'Source batch changed; review mappings before importing'
    book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    sheet = book['Sheet1']
    assert list(next(sheet.iter_rows(values_only=True))) == ['FT Material', 'Target Word', 'Target Phrase', 'Target Sentence'], 'Unexpected source headers'
    assert sheet.max_row == 156, 'This import is reviewed for the supplied 156-row batch only'
    for cell, (original, _, _) in CORRECTIONS.items():
        assert clean(sheet[cell].value) == original, f'Changed source at {cell}; review before reimport'
    catalog = {'word': {}, 'sentence': {}}
    deferred, changes, rows = [], [], []
    material = ''
    counts = {'materials': 0, 'wordRows': 0, 'phraseRows': 0, 'sentenceRows': 0}
    for row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        if values[0]:
            material = clean(values[0])
            counts['materials'] += 1
        rows.append({'row': row, 'material': material, 'word': values[1], 'phrase': values[2], 'sentence': values[3]})
        for column, kind, field in [(1, 'word', 'wordRows'), (2, 'phrase', 'phraseRows'), (3, 'sentence', 'sentenceRows')]:
            raw = values[column]
            if not raw:
                continue
            counts[field] += 1
            text = clean(raw)
            cell = f'{chr(65 + column)}{row}'
            origin = {'cell': cell, 'material': material, 'raw': raw}
            if kind == 'phrase' or (kind == 'word' and ' ' in text):
                deferred.append({**origin, 'reason': '短语暂不接入；包括单词列中的多词表达'})
                continue
            if '…' in text or '..' in text:
                deferred.append({**origin, 'reason': '未完成句型，需要业务补齐朗读文本'})
                continue
            if cell in ('D80', 'D100'):
                deferred.append({**origin, 'reason': '需确认斜杠范围或 burger 的单复数表达，暂不猜测标准朗读文本'})
                continue
            if cell in CORRECTIONS:
                text = CORRECTIONS[cell][1]
                changes.append({**origin, 'targets': [text], 'reason': CORRECTIONS[cell][2]})
            if kind == 'word' and text.endswith('/s'):
                base = text[:-2]
                targets = [base, base + 's']
            elif kind == 'sentence' and '/' in text:
                assert row in EXPANSIONS, f'Unreviewed slash notation at {cell}'
                targets = EXPANSIONS[row]
            else:
                targets = [text]
            if len(targets) > 1:
                changes.append({**origin, 'targets': targets, 'reason': '展开原表中明确的单复数或斜杠选项'})
            for target in targets:
                if kind == 'sentence' and target[-1] not in '.?!':
                    target += '.'
                assert '/' not in target and '…' not in target
                key = target.casefold()
                if key not in catalog[kind]:
                    ident = 'ft-word-' + target.lower() if kind == 'word' else 'ft-sentence-' + hashlib.sha256(target.encode()).hexdigest()[:12]
                    catalog[kind][key] = {'id': ident, 'text': target, 'kind': kind, 'sources': []}
                catalog[kind][key]['sources'].append(origin)
    items = [*catalog['word'].values(), *catalog['sentence'].values()]
    content = [{'id': item['id'], 'text': item['text'], 'ipa': '', 'zh': '', 'kind': '关键词' if item['kind'] == 'word' else '关键句', 'hint': '', 'example': ''} for item in items]
    counts.update(words=len(catalog['word']), sentences=len(catalog['sentence']))
    file = ROOT / 'src/features/practice/lib/catalog.ts'
    grammar = 'export const grammar = ' + file.read_text().split('export const grammar = ', 1)[1]
    file.write_text(f"// Imported from FT Target Words.xlsx; see docs/content/ft-target-words-2026-09-14.json.\nexport const CONTENT_VERSION = '{VERSION}';\nexport const DEMO_THRESHOLD = 80;\nexport const words = " + json.dumps(content, ensure_ascii=False, indent=2) + ' as const;\n' + grammar)
    folder = ROOT / 'docs/content'
    folder.mkdir(parents=True, exist_ok=True)
    manifest = {'sourceFile': source.name, 'sourceSha256': hashlib.sha256(source.read_bytes()).hexdigest(), 'sheet': 'Sheet1', 'contentVersion': VERSION, 'counts': counts, 'policy': 'B 列单词与 D 列完整句子；重复合并、明确选项展开；C 列与 B 列多词短语、未完成句型暂缓。原文件不修改。', 'items': items, 'deferred': deferred, 'normalizations': changes, 'sourceRows': rows}
    (folder / 'ft-target-words-2026-09-14.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps(counts, ensure_ascii=False))


if __name__ == '__main__':
    main(Path(sys.argv[1]))
